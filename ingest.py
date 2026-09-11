"""Load CSV/Excel into sqlite and write a copy under data/."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from db import ROOT, upsert_game
from spreads import norm_team, to_home_margin

PICKS_DIR = ROOT / "data" / "picks"
VEGAS_DIR = ROOT / "data" / "vegas"


def _rows_from_csv(text: str) -> list[dict]:
    f = io.StringIO(text)
    reader = csv.DictReader(f)
    rows = []
    for raw in reader:
        row = { (k or "").strip(): (v or "").strip() for k, v in raw.items() if k and k.strip() }
        if any(row.values()):
            rows.append(row)
    return rows


def _rows_from_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(it)]
    rows = []
    for vals in it:
        row = {}
        for h, v in zip(header, vals):
            if not h:
                continue
            row[h] = "" if v is None else str(v).strip()
        if any(row.values()):
            rows.append(row)
    return rows


def parse_upload(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return _rows_from_xlsx(data)
    return _rows_from_csv(data.decode("utf-8-sig"))


def _spread_text(row: dict) -> str:
    for key in ("spread", "fair_spread", "line", "vegas"):
        if row.get(key):
            return row[key]
    raise ValueError(f"no spread column in {list(row)}")


def ingest_rows(conn, rows: list[dict], *, season: int, week: int, kind: str) -> int:
    if kind not in ("model", "vegas"):
        raise ValueError("kind must be model or vegas")
    n = 0
    for row in rows:
        away = norm_team(row.get("away_team") or row.get("away") or "")
        home = norm_team(row.get("home_team") or row.get("home") or "")
        if not away or not home:
            raise ValueError(f"missing teams in {row}")
        margin = to_home_margin(away, home, _spread_text(row))
        kw = {"model_margin": margin} if kind == "model" else {"vegas_margin": margin}
        upsert_game(conn, season=season, week=week, away_team=away, home_team=home, **kw)
        n += 1
    conn.commit()
    return n


def save_upload(kind: str, week: int, filename: str, data: bytes) -> Path:
    dest_dir = PICKS_DIR if kind == "model" else VEGAS_DIR
    dest_dir.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename).suffix.lower() or ".csv"
    if suffix not in {".csv", ".xlsx", ".xls"}:
        suffix = ".csv"
    path = dest_dir / f"week_{week:02d}{suffix}"
    path.write_bytes(data)
    return path


def seed_week01(conn) -> int:
    path = PICKS_DIR / "week_01.csv"
    if not path.is_file():
        return 0
    rows = parse_upload(path.name, path.read_bytes())
    return ingest_rows(conn, rows, season=2026, week=1, kind="model")
